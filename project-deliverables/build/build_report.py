"""Build the CodeBridge project report (DOCX) with python-docx.

Every number and feature here comes from the repository (source, tests,
docs/TEST_MATRIX.md) or from verification runs recorded in CLAUDE.md.
Run: python build_report.py  -> ../CodeBridge_Project_Report.docx
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "assets"
OUT = ROOT / "CodeBridge_Project_Report.docx"

NAVY = RGBColor(0x0F, 0x2A, 0x4A)
TEAL = RGBColor(0x0D, 0x94, 0x88)
GREY = RGBColor(0x47, 0x55, 0x69)
BODY_FONT = "Calibri"
HEAD_FONT = "Cambria"
CODE_FONT = "Consolas"
BODY_PT = 11

doc = Document()

# ------------------------------------------------------------------ page setup
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(2.0)
sec.header_distance = Cm(1.0)
sec.footer_distance = Cm(1.0)
sec.different_first_page_header_footer = True
TEXT_WIDTH = Cm(21.0 - 4.4)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = BODY_FONT
normal.font.size = Pt(BODY_PT)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.12

for name, size, color, before, after in [("Heading 1", 14, NAVY, 10, 4), ("Heading 2", 11.5, TEAL, 7, 2)]:
    st = styles[name]
    st.font.name = HEAD_FONT
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    st.element.rPr.rFonts.set(qn("w:asciiTheme"), "")  # use the explicit font
    rf = st.element.rPr.rFonts
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(attr), HEAD_FONT)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

cap = styles["Caption"]
cap.font.name = BODY_FONT
cap.font.size = Pt(9)
cap.font.italic = True
cap.font.bold = False
cap.font.color.rgb = GREY
cap.paragraph_format.space_after = Pt(8)
cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

for lst in ("List Bullet", "List Number"):
    styles[lst].font.size = Pt(BODY_PT)
    styles[lst].paragraph_format.space_after = Pt(1.5)


# ------------------------------------------------------------------ helpers
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def add_runs(p, text, size=None, color=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = CODE_FONT
            r.font.size = Pt((size or BODY_PT) - 1)
        else:
            r = p.add_run(part)
        if size and not part.startswith("`"):
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p


def para(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=None, size=None):
    p = doc.add_paragraph()
    p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return add_runs(p, text, size=size)


def h1(text):
    return doc.add_heading(text, level=1)


def h2(text):
    return doc.add_heading(text, level=2)


def bullets(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_runs(p, it)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


TAB = [0]


def table(headers, rows, widths, size=9, caption=None):
    if caption:
        TAB[0] += 1
        caption = f"Table {TAB[0]}. {caption}"
        c = doc.add_paragraph(style="Caption")
        c.paragraph_format.keep_with_next = True
        c.paragraph_format.space_after = Pt(3)
        c.add_run(caption)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_margins(t)
    for i, htext in enumerate(headers):
        cell = t.rows[0].cells[i]
        shade(cell, "0F2A4A")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(htext)
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, text in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_runs(p, text, size=size)
            if ri % 2 == 1:
                shade(cells[i], "F1F5F9")
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    # keep header row repeated and rows unsplit
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for row in t.rows:
        cant = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = 0.6
    return t


FIG = [0]


def figure(path, caption, width_cm):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    FIG[0] += 1
    c = doc.add_paragraph(style="Caption")
    c.add_run(f"Figure {FIG[0]}. {caption}")


def figures_side_by_side(items, widths_cm):
    """Images in a borderless table, each with its own caption underneath."""
    if not isinstance(widths_cm, (list, tuple)):
        widths_cm = [widths_cm] * len(items)
    t = doc.add_table(rows=2, cols=len(items))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, (path, caption) in enumerate(items):
        for row in t.rows:
            row.cells[i].width = Cm(widths_cm[i] + 0.3)
        p = t.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        p.add_run().add_picture(str(path), width=Cm(widths_cm[i]))
        FIG[0] += 1
        c = t.rows[1].cells[i].paragraphs[0]
        c.style = doc.styles["Caption"]
        c.add_run(f"Figure {FIG[0]}. {caption}")
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def code(lines, size=8.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.left_indent = Cm(0.3)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F1F5F9")
    pPr.append(shd)
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.font.name = CODE_FONT
        r.font.size = Pt(size)
        if i < len(lines) - 1:
            r.add_break()
    return p


def code_pair(left_title, left, right_title, right, size=8.3, widths=(8.3, 8.3)):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, (title, lines) in enumerate(((left_title, left), (right_title, right))):
        cell = t.rows[0].cells[i]
        cell.width = Cm(widths[i])
        shade(cell, "F1F5F9")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = TEAL
        p2 = cell.add_paragraph()
        p2.paragraph_format.line_spacing = 1.0
        p2.paragraph_format.space_after = Pt(2)
        for j, line in enumerate(lines):
            rr = p2.add_run(line)
            rr.font.name = CODE_FONT
            rr.font.size = Pt(size)
            if j < len(lines) - 1:
                rr.add_break()
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = 0.6


def field(run, instr):
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)
        elif text == instr:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instr} "
            run._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)


# ------------------------------------------------------------------ footer/header
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for text, instr in (("CodeBridge: Project Report   ·   Page ", None), (None, "PAGE"), (" of ", None), (None, "NUMPAGES")):
    r = footer_p.add_run(text or "")
    r.font.size = Pt(8.5)
    r.font.color.rgb = GREY
    if instr:
        field(r, instr)
header_p = sec.header.paragraphs[0]
header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = header_p.add_run("Compiler Design Mini-Project  ·  CodeBridge")
hr.font.size = Pt(8.5)
hr.font.color.rgb = GREY

# ================================================================== COVER (page 1)
for _ in range(3):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("COMPILER DESIGN MINI-PROJECT REPORT")
r.font.size = Pt(11)
r.font.color.rgb = TEAL
r.bold = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(18)
r = p.add_run("CodeBridge")
r.font.name = HEAD_FONT
r.font.size = Pt(40)
r.bold = True
r.font.color.rgb = NAVY
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("A Pseudocode-to-C++ Compiler with a Phase-by-Phase Web Workbench")
r.font.name = HEAD_FONT
r.font.size = Pt(16)
r.font.color.rgb = NAVY
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
r = p.add_run("Lexer · Recursive-descent parser · Semantic analysis · Three-address code · Optimizer · C++17 code generation")
r.font.size = Pt(10)
r.font.color.rgb = GREY

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(26)
p.add_run().add_picture(str(ASSETS / "screenshots" / "live-cpp-fizzbuzz.png"), width=Cm(13.5))

doc.add_paragraph().paragraph_format.space_after = Pt(16)
cover = doc.add_table(rows=0, cols=2)
cover.alignment = WD_TABLE_ALIGNMENT.CENTER
for label, value in [
    ("Submitted by", "[Student Name]"),
    ("Roll number", "[Roll Number]"),
    ("Department", "[Department]"),
    ("Institution", "[Institution]"),
    ("Project guide", "[Guide Name]"),
    ("Academic year", "[Academic Year]"),
    ("Live application", "https://codebridge-swart.vercel.app"),
    ("Source code", "https://github.com/ashmita-dev/CodeBridge"),
]:
    cells = cover.add_row().cells
    cells[0].width, cells[1].width = Cm(4.2), Cm(9.0)
    a = cells[0].paragraphs[0]
    a.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    a.paragraph_format.space_after = Pt(2)
    ra = a.add_run(label + ":")
    ra.bold = True
    ra.font.color.rgb = NAVY
    b = cells[1].paragraphs[0]
    b.paragraph_format.space_after = Pt(2)
    b.add_run(value)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ================================================================== ABSTRACT
h1("Abstract")
para("CodeBridge is a compiler that translates a small, statically typed pseudocode language into C++17 and exposes "
     "every compiler phase in a web workbench. The compiler is written by hand in Python, without lexer or parser "
     "generators. It consists of a scanner, a recursive-descent parser that builds a typed abstract syntax tree (AST), "
     "a semantic analyser with a scoped symbol table and type checker, a three-address code (TAC) generator, an "
     "optimizer that applies eleven classical passes until a fixed point, and a code generator that rebuilds "
     "structured C++ from the optimized TAC. A FastAPI service exposes the pipeline through `POST /compile`, and a "
     "React and CodeMirror front end shows the tokens, AST, symbol table, original and optimized TAC, generated C++ "
     "and structured diagnostics for the program being edited. The system is deployed with the front end on Vercel "
     "and the API on Render. Correctness is checked by 213 automated backend tests and 20 frontend tests. The "
     "generated C++ of every example program is compiled with g++ and its output compared with an independent "
     "reference interpreter, and the optimizer is checked against 300 randomly generated programs.")
p = para("**Keywords:** compiler design, lexical analysis, recursive-descent parsing, abstract syntax tree, symbol table, "
         "type checking, three-address code, code optimization, source-to-source translation, C++, FastAPI, React.",
         align=WD_ALIGN_PARAGRAPH.LEFT)

# ================================================================== 1 INTRODUCTION
h1("1. Introduction")
para("A compiler course usually teaches lexical analysis, parsing, semantic analysis, intermediate code and "
     "optimization as separate topics, often on paper. Students rarely see one program travel through a complete, "
     "working compiler or observe how an optimizer rewrites intermediate code. CodeBridge was built to close that gap: "
     "it is a complete source-to-source compiler whose every intermediate artifact is visible and interactive in the "
     "browser, while the output is a real, compilable C++ program.")
h2("1.1 Problem statement")
para("Design and implement a compiler that accepts programs in a simple pseudocode language, reports lexical, syntax "
     "and semantic errors precisely, generates and optimizes intermediate code, and emits equivalent C++ source. The "
     "result of every phase must be presented to the user through a web interface, and the translation must be "
     "demonstrably correct: the C++ program must print exactly what the source program means.")
h2("1.2 Motivation")
bullets([
    "**Learning by observation:** seeing tokens, the AST, the symbol table and TAC for one's own program makes the "
    "theory concrete.",
    "**Correctness as a requirement:** a translator that produces plausible but wrong C++ is worse than none, so the "
    "project emphasizes verified semantics over feature count.",
    "**Accessibility:** a hosted web application needs no local toolchain, so anyone can try the compiler from a browser.",
])

# ================================================================== 2 OBJECTIVES & SCOPE
h1("2. Objectives and Scope")
h2("2.1 Objectives")
bullets([
    "Design a small but complete language: variables, the types int, float, string and bool, arithmetic, "
    "comparisons, logical operators, `print`, `if`/`else if`/`else` and `while`.",
    "Implement every compiler phase by hand and keep each phase testable without the web framework.",
    "Report errors as structured diagnostics with a stage, a code, a message and an exact line and column.",
    "Generate three-address code, optimize it safely, and generate C++17 that behaves exactly like the source.",
    "Present all phases in a responsive web workbench and deploy it publicly.",
], style="List Number")
h2("2.2 Scope")
para("**In scope:** the language above; all six compiler phases; a JSON API; the web workbench; automated and live "
     "verification. **Out of scope by design:** functions, arrays, `for` loops, `break`/`continue` and input; and "
     "executing user programs on the server. CodeBridge generates C++ but never runs it in production, which keeps "
     "the hosted service free of arbitrary code execution.")

# ================================================================== 3 REQUIREMENTS
h1("3. Requirements")
table(["ID", "Requirement", "How it is met"], [
    ["F1", "Translate valid programs to compilable C++17", "Six-phase pipeline; g++ build-and-run tests"],
    ["F2", "Show tokens, AST, symbols, TAC, optimized TAC and C++", "API returns every artifact; one tab per artifact"],
    ["F3", "Report lexical, syntax and semantic errors with line:column", "Structured diagnostics (codes L, P, S)"],
    ["F4", "No downstream artifacts after a failing phase", "Pipeline stops; TAC and C++ are null on failure"],
    ["F5", "Optimize intermediate code without changing behaviour", "Eleven passes; differential testing"],
    ["N1", "User input must never crash the server", "Input limits, recursion limits, JSON 500 handler"],
    ["N2", "Never execute user-generated C++ on the backend", "No execution endpoint exists"],
    ["N3", "Configuration through environment variables, no secrets", "`CORS_ORIGINS`, `MAX_SOURCE_CHARS`, `VITE_API_URL`"],
    ["N4", "Responsive, accessible interface", "Responsive layout down to phone widths (Figure 9)"],
], [1.1, 7.6, 7.9], caption="Functional (F) and non-functional (N) requirements")

# ================================================================== 4 TECH STACK
h1("4. Technology Stack")
table(["Layer", "Technology (version in the repository)", "Role"], [
    ["Compiler and API", "Python 3.13.15, FastAPI 0.143.0, Pydantic 2.14.1, Uvicorn 0.54.0",
     "Compiler package, HTTP API, validation"],
    ["Frontend", "React 19.3, TypeScript 7.0, Vite 8.3, Tailwind CSS 4.3", "Workbench UI and build"],
    ["Editor", "CodeMirror 6 (lang-cpp, lint, language)", "Source editor, C++ viewer, error underlines"],
    ["Testing", "pytest 9.1, httpx2, Vitest 5.0, g++ 16.1 (MSYS2)", "Unit, API, end-to-end and build-and-run tests"],
    ["CI and hosting", "GitHub Actions, Render (API), Vercel (frontend)", "Automated tests; production deployment"],
], [3.0, 8.1, 5.5], caption="Technology stack")

# ================================================================== 5 ARCHITECTURE
h1("5. System Architecture")
para("CodeBridge is a two-tier web application (Figure 1). The browser runs a single-page React application served by "
     "Vercel. It sends the editor contents to the FastAPI backend on Render and renders the JSON response. The "
     "compiler lives in the package `app/compiler`, which has no web dependencies; `pipeline.compile_source()` runs "
     "the phases in order, records the time each one takes and stops at the first phase that reports errors.")
figure(ASSETS / "diagrams" / "architecture.png", "System architecture, drawn from the modules in frontend/src and backend/app.", 14.6)
h2("5.1 API integration")
para("The API has two endpoints. `GET /health` returns `{\"status\": \"ok\"}` and is used by Render's health check "
     "and by the frontend's online indicator. `POST /compile` accepts `{\"source\": ..., \"optimize\": true}` and "
     "returns `success`, the status of each phase (`ok`, `error` or `skipped`), the diagnostics, the tokens, the AST, "
     "the symbol table, the original and optimized TAC, the C++ source, optimization statistics and per-phase "
     "timings. Requests over `MAX_SOURCE_CHARS` (20,000) are rejected with HTTP 413, malformed bodies with 422, and an "
     "unexpected exception produces a JSON 500 response instead of a crash. CORS admits only the origins listed in "
     "`CORS_ORIGINS`; in production that is exactly `https://codebridge-swart.vercel.app`. The frontend's API base "
     "URL is inlined at build time from `VITE_API_URL`. In auto mode it recompiles 650 ms after the last keystroke, "
     "and on a slow first request it tells the user the free-tier server is waking up.")
h2("5.2 Deployment")
table(["Component", "Platform and configuration"], [
    ["Backend API", "Render web service from the `render.yaml` Blueprint: root `backend`, `pip install -r requirements.txt`, "
                    "`uvicorn app.main:app --host 0.0.0.0 --port $PORT`, health check `/health`, `PYTHON_VERSION=3.13.15`, "
                    "`CORS_ORIGINS` set to the Vercel origin. URL: https://codebridge-api.onrender.com"],
    ["Frontend", "Vercel project built by `vercel.json` (`npm ci` and `npm run build` in `frontend/`, output "
                 "`frontend/dist`), `VITE_API_URL=https://codebridge-api.onrender.com`. URL: https://codebridge-swart.vercel.app"],
    ["Continuous delivery", "Both platforms redeploy on every push to `main`; GitHub Actions runs all tests on each push."],
], [3.2, 13.4], caption="Production deployment")

# ================================================================== 6 METHODOLOGY
h1("6. Methodology")
para("The compiler follows the classical analysis–synthesis structure (Figure 2). Each phase is a separate module with "
     "its own tests, and each consumes only the output of the previous phase. Development was incremental and "
     "test-driven: a phase was added only when the previous one was complete and tested, in the order language "
     "specification, lexer, parser, semantic analysis, TAC, optimizer and code generator, then API, frontend and "
     "deployment. To check meaning rather than just form, the test suite contains two **reference interpreters**, one "
     "for the AST and one for TAC. They are used as oracles for the optimizer and the code generator.")
figure(ASSETS / "diagrams" / "pipeline.png", "The six-phase compiler pipeline and its error paths.", 15.6)

# ================================================================== 7 COMPILER DESIGN
h1("7. Compiler Design")
h2("7.1 Source language")
para("Statements end with a newline or `;`, blocks use braces and `#` starts a comment. A declaration is "
     "`let x = expr` or `let x: type = expr`; the only implicit conversion is int to float. `print` takes one or more "
     "comma-separated expressions, with or without parentheses (`print(a, b)` or `print a, b`). Semantics follow C++ "
     "deliberately: integers are 64-bit, `/` truncates toward zero, `%` takes the sign of the dividend, `and`/`or` "
     "short-circuit and floats print like `std::cout`. The grammar in EBNF (abridged from `parser.py`):")
code([
    'statement  = let_stmt | assign_stmt | print_stmt | if_stmt | while_stmt ;',
    'print_stmt = "print" "(" [ args ] ")" | "print" args ;',
    'if_stmt    = "if" expression block [ "else" ( if_stmt | block ) ] ;',
    'expression = or_expr ;   or_expr = and_expr { "or" and_expr } ;   ...',
    'term       = unary { ( "*" | "/" | "%" ) unary } ;   unary = "-" unary | primary ;',
])
h2("7.2 Lexical analysis")
para("`lexer.py` is a hand-written scanner over the character stream. Operators are matched longest-first (`<=` "
     "before `<`), every token carries a 1-based span, and newlines become `NEWLINE` tokens except inside parentheses, "
     "so expressions may span lines. The lexer records an error and continues, so one pass reports every lexical "
     "error: invalid characters with hints such as \"use `not`\" for `!` (L001), malformed numbers such as `1.2.3` "
     "(L002), literals outside the 64-bit or finite float range (L003), unterminated strings (L004) and unknown "
     "escape sequences (L005).")
h2("7.3 Syntax analysis and the AST")
para("`parser.py` is a predictive recursive-descent parser with one method per non-terminal. Operator precedence "
     "comes from the chain of grammar levels (`or` < `and` < `not` < equality < comparison < additive < "
     "multiplicative < unary), and left associativity from iteration inside each level. The parser builds dataclass "
     "nodes (`ast_nodes.py`) that carry source spans and serialize to JSON for the AST view. It stops at the first "
     "syntax error with a precise message, for example \"expected ')' to close '(' but found '{'\" (P001). Nesting depth "
     "(50) and expression height (150) are bounded (P002), so hostile input cannot overflow the Python stack.")
h2("7.4 Semantic analysis and the symbol table")
para("`semantic.py` walks the AST with a symbol table implemented as a stack of scopes, one per block. It resolves "
     "names (undefined or undeclared variables, S001; redeclaration in one scope, S002), infers the type of every "
     "expression and checks initializers and assignments (S003), operand types (S004), `bool` conditions (S005) and "
     "division by a literal zero (S006). All semantic errors are collected; a variable with an invalid initializer is "
     "marked so its later uses do not cause cascading errors. Shadowed variables receive unique names (`x`, `x_2`) "
     "that also avoid C++ keywords and the TAC temporaries `tN`. Unused variables produce warning S100.")
h2("7.5 Intermediate code: three-address code")
para("`tacgen.py` lowers the typed AST into typed TAC instructions: copies, binary and unary operations (including an "
     "explicit `(float)` conversion), labels, `goto`, `ifFalse`/`ifTrue` conditional jumps and `print`. Control flow "
     "uses the standard schemes, and `and`/`or` are compiled with jumps so they short-circuit; for example "
     "`d != 0 and 10 / d > 1` never divides by zero.")
table(["Instruction", "Example", "Meaning"], [
    ["Copy", "`x = t1`", "assign a value or temporary to a variable"],
    ["Binary / unary", "`t1 = a * b`, `t2 = -a`, `t3 = not c`", "one operator, at most two operands"],
    ["Conversion", "`t4 = (float) i`", "explicit int-to-float widening"],
    ["Jumps", "`goto L1`, `ifFalse t1 goto L2`, `ifTrue t1 goto L3`", "unconditional and conditional control flow"],
    ["Label / output", "`L1:`, `print a, b`", "jump target; print values separated by spaces"],
], [3.0, 7.0, 6.6], caption="Three-address code instruction forms (tac.py)")
h2("7.6 Optimization")
para("`optimizer.py` repeats its passes until nothing changes. Constant and copy propagation, constant folding and "
     "algebraic simplification are local to basic blocks: known values are discarded at every label, where control "
     "paths merge, and whenever a variable they depend on is reassigned. Folding uses C++ semantics and refuses "
     "overflow, division by zero and non-finite results; floating-point identities are applied only when exact under "
     "IEEE-754. The other passes are self-assignment removal, branch folding, temporary coalescing, dead-temporary "
     "elimination, unreachable-code removal on the control-flow graph, and redundant-jump and unused-label removal. "
     "User variables and `print` statements are never removed. The example below is the exact compiler output for "
     "`let x = 2 * 3 + 4` followed by a loop; on the bundled optimizer showcase program the instruction count falls "
     "from 25 to 14 (−44%, Figure 6).")
code_pair("Original TAC", ["    t1 = 2 * 3", "    t2 = t1 + 4", "    x = t2", "    i = 0", "L1:",
                           "    t3 = i < x", "    ifFalse t3 goto L2"],
          "Optimized TAC", ["    x = 10", "    i = 0", "L1:", "    t3 = i < x", "    ifFalse t3 goto L2", "", ""])
h2("7.7 C++ code generation")
para("`codegen_cpp.py` translates the optimized TAC, never the AST, so the C++ reflects the optimizations. It works in "
     "three steps: temporaries used once are folded back into expressions; jump patterns are recognized and rebuilt "
     "as `while`, `if`/`else` and `else if` (with a `goto` fallback that keeps the output correct if a pattern does "
     "not match); and the program is emitted with int as `long long`, float as `double`, string as `std::string`, "
     "parentheses chosen by precedence, escaped string literals and `std::cout` with `std::boolalpha`. The pair below "
     "is unedited compiler output; built with g++, the C++ prints `1 is odd`, `3 is odd`, `5 is odd` and "
     "`sum of even squares: 20`, the same as the reference interpreter.")
code_pair("CodeBridge source", [
    "let n = 1", "let total = 0", "while n <= 5 {", "    if n % 2 == 0 {", "        total = total + n * n",
    "    } else {", '        print n, "is odd"', "    }", "    n = n + 1", "}",
    'print("sum of even squares:", total)',
], "Generated C++ (body of main)", [
    "long long n = 1;", "long long total = 0;", "while (n <= 5) {", "    if (n % 2 == 0) {",
    "        total = total + n * n;", "    } else {", '        std::cout << n << " is odd\\n";', "    }",
    "    n = n + 1;", "}", 'std::cout << "sum of even squares: " << total << "\\n";',
], widths=(6.6, 10.0))
h2("7.8 Error handling")
para("Every diagnostic has a stage (lexical, syntax, semantic or internal), a severity, a code, a message, a start "
     "and end position and optional hints. The editor underlines each error at its span, and clicking a diagnostic "
     "moves the cursor to it (Figure 8). When any phase fails, later phases are reported as skipped and the response "
     "contains no TAC and no C++, so the interface never shows stale or invented output.")

# ================================================================== 8 IMPLEMENTATION & SCREENSHOTS
h1("8. Implementation and Screenshots")
para("The compiler package is about 2,600 lines of Python in fourteen files, the FastAPI layer about 150 lines and "
     "the backend tests about 1,500 lines. The frontend is about 1,900 lines of TypeScript and TSX: a header with the online indicator, a pipeline "
     "strip with per-phase status and timing, the editor panel with an example selector, reset, optimize and auto-"
     "compile toggles, and an output panel with seven tabs and copy buttons. Deep links such as "
     "`?example=fizzbuzz&tab=cpp` open a given example and tab. All screenshots below were captured from the live "
     "deployment at https://codebridge-swart.vercel.app on 13 October 2026.")
figure(ASSETS / "screenshots" / "live-cpp-fizzbuzz.png",
       "Live workbench: FizzBuzz source (left) and the generated C++ with a recovered else-if chain (right).", 15.6)
CROPS = ASSETS / "screenshots" / "crops"
figures_side_by_side([
    (CROPS / "tokens-panel.png", "Tokens tab: type, lexeme, value and line:column, with category filters (cropped)."),
    (CROPS / "ast-panel.png", "AST tab: typed nodes of an else-if chain; clicking a node selects its source (cropped)."),
], 7.9)
figures_side_by_side([
    (CROPS / "optimized-panel.png", "Optimized TAC with per-pass statistics: 25 → 14 instructions (cropped)."),
    (CROPS / "symbols-panel.png", "Symbol table: types, scope depth and the renamed shadowed variable x_2 (cropped)."),
], 7.9)
figures_side_by_side([
    (CROPS / "diagnostics-panel.png", "Diagnostics tab: five semantic errors with codes, line:column positions and hints; later phases not run (cropped)."),
    (CROPS / "mobile-top.png", "Phone-width layout, 500 px viewport (cropped)."),
], [11.4, 4.2])

# ================================================================== 9 TESTING
h1("9. Testing Methodology")
para("Testing combines automated suites, run locally and by GitHub Actions on every push, with manual and scripted "
     "checks of the deployed system. Each example program is checked four ways and the outputs must be identical: the "
     "AST interpreter, a TAC interpreter on the original TAC, the same interpreter on the optimized TAC, and a native "
     "binary built from the generated C++ with g++ (`-std=c++17 -Wall -Wextra`). The optimizer is additionally "
     "checked against 300 random programs, and nine programs with hand-derived expected output guard against a "
     "mistake shared by the interpreter and the compiler.")
table(["Suite (file)", "Tests", "What it covers"], [
    ["Lexer (`test_lexer.py`)", "12", "token kinds, longest match, positions, CRLF, every lexical error code"],
    ["Parser and AST (`test_parser.py`)", "48", "15 precedence cases, statements, both print forms, 21 syntax-error cases"],
    ["Semantic (`test_semantic.py`)", "31", "S001–S006 with exact positions, scopes, shadowing, no cascades"],
    ["TAC (`test_tac.py`)", "18", "exact TAC shapes, short-circuit, TAC vs. AST interpreter"],
    ["Optimizer (`test_optimizer.py`)", "28", "safe folding, no propagation into loops, 300 random programs"],
    ["C++ generation (`test_codegen.py`)", "36", "structure, parentheses, escaping; g++ build and run of all programs"],
    ["API and health (`test_api.py`, `test_health.py`)", "34", "payloads, 413/422/500, CORS, every example over HTTP"],
    ["End-to-end (`test_e2e_server.py`)", "6", "live Uvicorn server → HTTP → g++ → output"],
    ["Frontend (Vitest, 6 files)", "20", "API client, AST view, phases, examples, positions, output panel"],
], [5.6, 1.3, 9.7], caption="Automated test suites (213 backend + 20 frontend)")

# ================================================================== 10 RESULTS
h1("10. Verified Results")
h2("10.1 Automated results")
para("On the final commit all **213 backend tests passed**, including the g++ build-and-run tests and the end-to-end "
     "server test, and the frontend type check, **20 Vitest tests** and the production build passed. GitHub Actions CI "
     "(Ubuntu) passed on the same commits. Locally, a 19,860-character program compiled through all six phases in "
     "0.057 s (median of five runs on the Windows 11 development machine).")
para("Table 6 lists the twelve bundled example programs. Every valid example is compiled, built with g++ and run by "
     "the test suite, and its output must equal the stored expected output; every error example must fail in the "
     "intended phase with no TAC or C++. The TAC counts are the compiler's own optimization statistics. Programs whose "
     "branch conditions depend on constants shrink the most (`conditions`: 43 to 7), while loop-heavy programs shrink "
     "little because propagation stops at loop labels.")
table(["Example", "Exercises", "TAC before → after", "Verified result"], [
    ["hello", "variables, strings, print", "10 → 9", "g++ output = expected"],
    ["arithmetic", "precedence, associativity, int/float division", "36 → 14", "g++ output = expected"],
    ["conditions", "else-if chains, nested branches, and/or/not", "43 → 7", "g++ output = expected"],
    ["loops", "sum, factorial, countdown", "35 → 30", "g++ output = expected"],
    ["fizzbuzz", "loop with an else-if chain", "30 → 29", "g++ output = expected"],
    ["primes", "nested loops, short-circuit conditions", "33 → 30", "g++ output = expected"],
    ["fibonacci", "block-scoped variables in a loop", "15 → 13", "g++ output = expected"],
    ["types", "annotations, widening, escapes, shadowing", "31 → 14", "g++ output = expected"],
    ["optimization", "folding, propagation, simplification", "25 → 14", "g++ output = expected"],
    ["error_lexical", "invalid characters, malformed numbers", "not generated", "3 errors (L001, L002)"],
    ["error_syntax", "unbalanced parenthesis", "not generated", "1 error (P001)"],
    ["error_semantic", "undefined variables, type mismatches", "not generated", "5 errors (S001, S003–S005)"],
], [2.8, 6.4, 3.2, 4.2], size=8.8, caption="Example programs and their verified results")
h2("10.2 Manual and scripted checks of the deployed system (13 October 2026)")
table(["Check", "Result"], [
    ["`GET https://codebridge-api.onrender.com/health`", "HTTP 200, `{\"status\":\"ok\"}`"],
    ["CORS preflight from the Vercel origin / from another origin", "allowed / rejected with HTTP 400"],
    ["Live app in headless Microsoft Edge", "\"Compiler online\"; FizzBuzz and the optimizer showcase compiled through all six phases"],
    ["C++ returned by the live API for fizzbuzz, primes and types, built with g++", "output identical to `examples/expected`"],
    ["Live syntax and semantic error requests", "P001 and S001 diagnostics, no C++ returned"],
    ["Warm `POST /compile` round trip from the development machine (5 requests)", "0.55 s to 1.45 s"],
    ["Pre-deployment audit: 21 hand-checked programs, 18 error cases, 200 random programs", "all matched; 192 random programs matched g++, 8 skipped as integer-overflow (undefined behaviour)"],
], [8.6, 8.0], caption="Checks performed manually or with one-off scripts (not part of the automated suite)")

# ================================================================== 11 LIMITATIONS & FUTURE
h1("11. Limitations and Future Scope")
bullets([
    "**Language:** no functions, arrays, `for` loops, `break`/`continue` or input.",
    "**Error recovery:** the parser reports only the first syntax error (lexical and semantic errors are all reported).",
    "**Optimization scope:** propagation is local to basic blocks, so loop-invariant values are not propagated and "
    "there is no loop-invariant code motion.",
    "**Run-time behaviour:** division by a variable that is zero, or integer overflow, is undefined behaviour exactly as "
    "in C++; the generated program is not executed by the service.",
    "**Hosting:** the free Render instance sleeps when idle, so the first request can take up to about a minute.",
])
para("**Future scope:** functions with a call stack in TAC; panic-mode error recovery; global data-flow analysis "
     "(reaching definitions, liveness) and SSA form; and running the generated C++ in the browser with WebAssembly, "
     "which would show program output without executing code on the server.")

# ================================================================== 12 CONCLUSION
h1("12. Conclusion")
para("CodeBridge implements a complete compiler pipeline for a small language and makes each phase observable in a "
     "deployed web application. Hand-written lexing, recursive-descent parsing, scoped semantic analysis, TAC "
     "generation, safe optimization and structured C++ generation work together, and testing against reference "
     "interpreters and real g++ binaries, both automatically and on the live deployment, gives strong evidence that "
     "the generated C++ preserves the meaning of the source program.")

# ================================================================== REFERENCES
h1("References")
refs = [
    "A. V. Aho, M. S. Lam, R. Sethi and J. D. Ullman, Compilers: Principles, Techniques, and Tools, 2nd ed. Pearson, 2006.",
    "K. D. Cooper and L. Torczon, Engineering a Compiler, 2nd ed. Morgan Kaufmann, 2011.",
    "R. Nystrom, Crafting Interpreters. Genever Benning, 2021.",
    "ISO/IEC 14882:2017, Programming languages: C++.",
    "FastAPI documentation, https://fastapi.tiangolo.com; Pydantic documentation, https://docs.pydantic.dev.",
    "React documentation, https://react.dev; Vite documentation, https://vite.dev; CodeMirror 6, https://codemirror.net.",
    "pytest documentation, https://docs.pytest.org; Vitest documentation, https://vitest.dev.",
    "Render documentation (Blueprints), https://render.com/docs; Vercel documentation, https://vercel.com/docs.",
]
for i, ref in enumerate(refs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.first_line_indent = Cm(-0.7)
    add_runs(p, f"[{i}]  {ref}", size=9.5)

doc.core_properties.title = "CodeBridge: Project Report"
doc.core_properties.subject = "Pseudocode-to-C++ compiler"
doc.core_properties.author = "[Student Name]"
doc.save(OUT)
print("wrote", OUT)
