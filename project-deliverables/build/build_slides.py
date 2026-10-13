"""Build the 15-slide CodeBridge presentation (16:9) with python-pptx.

Diagrams, code panels and tables are native, editable PowerPoint shapes.
Screenshots are the genuine captures in ../assets/screenshots.
Run: python build_slides.py  -> ../CodeBridge_Presentation.pptx
"""
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / "assets" / "screenshots"
CROPS = SHOTS / "crops"
OUT = ROOT / "CodeBridge_Presentation.pptx"

# ------------------------------------------------------------------ theme
NAVY = RGBColor(0x0B, 0x12, 0x20)
NAVY2 = RGBColor(0x13, 0x1C, 0x2E)
INK = RGBColor(0x0F, 0x17, 0x2A)
SLATE = RGBColor(0x47, 0x55, 0x69)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
LINE = RGBColor(0xCB, 0xD5, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL = RGBColor(0x0D, 0x94, 0x88)
TEAL_L = RGBColor(0x2D, 0xD4, 0xBF)
TEAL_BG = RGBColor(0xCC, 0xFB, 0xF1)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
BLUE_BG = RGBColor(0xDB, 0xEA, 0xFE)
RED = RGBColor(0xB9, 0x1C, 0x1C)
RED_BG = RGBColor(0xFE, 0xE2, 0xE2)
AMBER = RGBColor(0xB4, 0x53, 0x09)
AMBER_BG = RGBColor(0xFE, 0xF3, 0xC7)
CODE_BG = RGBColor(0x0B, 0x12, 0x20)
FONT = "Segoe UI"
FONT_SB = "Segoe UI Semibold"
MONO = "Consolas"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
TOTAL = 15


# ------------------------------------------------------------------ primitives
def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None, line_w=1.0):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def text(slide, x, y, w, h, content, size=16, color=INK, bold=False, font=FONT, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1, after=4, shape=None):
    """content: str or list of paragraphs; a paragraph is str or list of (text, opts) runs."""
    tb = shape or slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        runs = para if isinstance(para, list) else [(para, {})]
        for rt, opts in runs:
            r = p.add_run()
            r.text = rt
            f = r.font
            f.name = opts.get("font", font)
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.italic = opts.get("italic", False)
            f.color.rgb = opts.get("color", color)
    return tb


def bullets(slide, x, y, w, h, items, size=16, color=INK, gap=6, marker="▸", marker_color=TEAL):
    paras = []
    for it in items:
        runs = [(f"{marker}\t", {"color": marker_color, "bold": True})]
        runs += rich(it, color)
        paras.append(runs)
    tb = text(slide, x, y, w, h, paras, size=size, color=color, after=gap)
    hang = str(int(Inches(0.28)))
    for p in tb.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", hang)
        pPr.set("indent", "-" + hang)
    return tb


RICH = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def rich(s, color=INK):
    out = []
    for part in RICH.split(s):
        if not part:
            continue
        if part.startswith("**"):
            out.append((part[2:-2], {"bold": True, "color": color}))
        elif part.startswith("`"):
            out.append((part[1:-1], {"font": MONO, "color": TEAL if color != WHITE else TEAL_L}))
        else:
            out.append((part, {"color": color}))
    return out


def line(slide, x1, y1, x2, y2, color=INK, width=1.5, arrow=True, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if arrow:
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
    if dash:
        pd = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
        ln.append(pd)
    return c


def box(slide, x, y, w, h, title, sub=None, fill=TEAL_BG, edge=TEAL, title_size=14, sub_size=11,
        title_color=INK, sub_color=SLATE, radius=0.12):
    s = rect(slide, x, y, w, h, fill=fill, line=edge, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    paras = [[(title, {"bold": True, "size": title_size, "color": title_color, "font": FONT_SB})]]
    if sub:
        for sl in sub.split("\n"):
            paras.append([(sl, {"size": sub_size, "color": sub_color})])
    text(slide, 0, 0, 0, 0, paras, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0, spacing=1.0, shape=s)
    s.text_frame.margin_left = s.text_frame.margin_right = Inches(0.06)
    return s


KW = {"let", "if", "else", "while", "print", "and", "or", "not", "true", "false", "int", "float", "string", "bool",
      "long", "double", "return", "std", "include"}
TOKEN = re.compile(r'("(?:[^"\\]|\\.)*"|#.*$|//.*$|\b\d+(?:\.\d+)?\b|\b[A-Za-z_]+\b|\s+|.)')
CODE_COLORS = {"kw": RGBColor(0xC0, 0x84, 0xFC), "str": RGBColor(0x86, 0xEF, 0xAC), "num": RGBColor(0xFD, 0xBA, 0x74),
               "com": RGBColor(0x64, 0x74, 0x8B), "txt": RGBColor(0xE2, 0xE8, 0xF0), "lbl": TEAL_L}


def code_runs(line_text):
    runs = []
    if re.match(r"^\s*L\d+:\s*$", line_text):
        return [(line_text, {"font": MONO, "color": CODE_COLORS["lbl"]})]
    for tok in TOKEN.findall(line_text):
        if tok.startswith('"'):
            c = "str"
        elif tok.startswith("#") or tok.startswith("//"):
            c = "com"
        elif re.fullmatch(r"\d+(?:\.\d+)?", tok):
            c = "num"
        elif tok in KW or tok in ("goto", "ifFalse", "ifTrue", "cout", "boolalpha"):
            c = "kw"
        else:
            c = "txt"
        runs.append((tok, {"font": MONO, "color": CODE_COLORS[c]}))
    return runs or [(" ", {"font": MONO})]


def code_panel(slide, x, y, w, h, lines, title=None, size=13):
    panel = rect(slide, x, y, w, h, fill=CODE_BG, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    top = y + 0.12
    if title:
        text(slide, x + 0.2, y + 0.1, w - 0.4, 0.32, title, size=11, color=TEAL_L, bold=True)
        top = y + 0.45
    text(slide, x + 0.2, top, w - 0.35, h - (top - y) - 0.08, [code_runs(l) for l in lines], size=size,
         font=MONO, spacing=1.05, after=0)
    return panel


def picture(slide, path, x, y, w=None, h=None, border=True):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)
    if border:
        pic.line.color.rgb = LINE
        pic.line.width = Pt(0.75)
    return pic


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def content_slide(n, title, kicker):
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    rect(s, 0, 0, 13.333, 0.08, fill=TEAL)
    text(s, 0.6, 0.32, 11.2, 0.35, kicker.upper(), size=12, color=TEAL, bold=True, font=FONT_SB)
    text(s, 0.6, 0.62, 12.1, 0.75, title, size=30, color=INK, bold=True, font=FONT_SB)
    # footer
    rect(s, 0.6, 7.05, 12.13, 0.012, fill=LINE)
    text(s, 0.6, 7.1, 8, 0.3, "CodeBridge  ·  Pseudocode-to-C++ Compiler", size=10, color=SLATE)
    text(s, 10.73, 7.1, 2.0, 0.3, f"{n} / {TOTAL}", size=10, color=SLATE, align=PP_ALIGN.RIGHT)
    return s


def dark_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY
    return s


def stat(slide, x, y, w, value, label, color=TEAL, value_size=34):
    rect(slide, x, y, w, 1.5, fill=LIGHT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    text(slide, x, y + 0.06, w, 0.65, value, size=value_size, color=color, bold=True, font=FONT_SB,
         align=PP_ALIGN.CENTER)
    text(slide, x + 0.1, y + 0.76, w - 0.2, 0.66, label, size=12, color=SLATE, align=PP_ALIGN.CENTER,
         spacing=1.0, after=0)


def table(slide, x, y, w, rows, col_w, size=12, header_fill=NAVY2, row_h=0.36):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    t = shape.table
    for i, cw in enumerate(col_w):
        t.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        t.rows[r].height = Inches(row_h)
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (LIGHT if r % 2 == 0 else WHITE)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            for rt, opts in rich(val, WHITE if r == 0 else INK):
                run = p.add_run()
                run.text = rt
                run.font.size = Pt(size)
                run.font.name = opts.get("font", FONT)
                run.font.bold = r == 0 or opts.get("bold", False)
                run.font.color.rgb = opts["color"]
    # turn off banding styles that would override fills
    tblPr = shape._element.graphic.graphicData.tbl.tblPr
    tblPr.set("bandRow", "0")
    tblPr.set("firstRow", "0")
    return shape


# ================================================================== 1 TITLE
s = dark_slide()
rect(s, 0, 0, 0.18, 7.5, fill=TEAL)
text(s, 0.75, 0.75, 6.5, 0.4, "COMPILER DESIGN MINI-PROJECT", size=14, color=TEAL_L, bold=True, font=FONT_SB)
text(s, 0.75, 1.25, 6.6, 1.1, "CodeBridge", size=60, color=WHITE, bold=True, font=FONT_SB)
text(s, 0.75, 2.35, 6.4, 1.0, "Pseudocode-to-C++ Compiler", size=30, color=WHITE, font=FONT)
text(s, 0.75, 3.25, 6.2, 0.9, "A hand-written six-phase compiler with a live web workbench that shows every phase.",
     size=16, color=RGBColor(0xCB, 0xD5, 0xE1))
info = [
    [("Presented by  ", {"color": TEAL_L, "bold": True}), ("[Student Name]  ·  [Roll Number]", {"color": WHITE})],
    [("Guide  ", {"color": TEAL_L, "bold": True}), ("[Guide Name]", {"color": WHITE})],
    [("Department  ", {"color": TEAL_L, "bold": True}), ("[Department], [Institution]  ·  [Academic Year]", {"color": WHITE})],
]
text(s, 0.75, 4.45, 6.4, 1.2, info, size=14, after=6)
text(s, 0.75, 5.9, 6.4, 0.7, [
    [("Live  ", {"color": TEAL_L, "bold": True}), ("codebridge-swart.vercel.app", {"color": WHITE})],
    [("Code  ", {"color": TEAL_L, "bold": True}), ("github.com/ashmita-dev/CodeBridge", {"color": WHITE})],
], size=14, after=4)
picture(s, SHOTS / "live-cpp-fizzbuzz.png", 7.45, 1.55, w=5.4, border=False)
rect(s, 7.45, 1.55, 5.4, 3.375, fill=None, line=RGBColor(0x33, 0x41, 0x55))
text(s, 7.45, 5.0, 5.4, 0.35, "Live deployment: FizzBuzz compiled to C++", size=11,
     color=RGBColor(0x94, 0xA3, 0xB8), align=PP_ALIGN.CENTER)
notes(s, "Introduce CodeBridge: a compiler written from scratch in Python that translates a small pseudocode "
         "language into C++17. The screenshot is the live deployment compiling FizzBuzz. Fill in name, roll number, "
         "guide and institution before presenting.")

# ================================================================== 2 INTRODUCTION
s = content_slide(2, "Introduction and project overview", "Introduction")
bullets(s, 0.6, 1.65, 6.6, 4.6, [
    "CodeBridge translates a small, statically typed **pseudocode language** into **C++17** source code.",
    "Every phase is **hand-written in Python**: lexer, recursive-descent parser, semantic analysis, TAC, "
    "optimizer and C++ generator. No parser generators.",
    "A **web workbench** shows tokens, AST, symbol table, original and optimized TAC, C++ and diagnostics "
    "for the code being edited.",
    "Deployed: React frontend on **Vercel**, FastAPI backend on **Render**.",
], size=18, gap=12)
stat(s, 7.7, 1.75, 2.45, "6", "compiler phases")
stat(s, 10.3, 1.75, 2.45, "233", "automated tests\n(213 backend + 20 frontend)", value_size=34)
stat(s, 7.7, 3.4, 2.45, "12", "bundled example\nprograms")
stat(s, 10.3, 3.4, 2.45, "300", "random programs in the\noptimizer differential test")
rect(s, 7.7, 5.1, 5.05, 1.6, fill=TEAL_BG, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
text(s, 7.9, 5.2, 4.7, 1.45, [
    [("Output is real C++.", {"bold": True, "color": TEAL})],
    [("Generated programs are built with g++ in the test suite and must print exactly what the source program means.",
      {"color": INK})],
], size=15, after=4)
notes(s, "Give the one-sentence summary, then the numbers: six phases, 233 automated tests, twelve examples and "
         "a 300-program differential test. Stress that correctness of the generated C++ is checked by compiling it.")

# ================================================================== 3 PROBLEM
s = content_slide(3, "Problem statement", "Why this project")
cards = [
    ("Phases learned in isolation", "Lexing, parsing, semantic analysis, IR and optimization are usually taught "
     "separately and on paper.", BLUE, BLUE_BG),
    ("Optimizers are invisible", "Students rarely see how an optimizer rewrites intermediate code, or why a "
     "transformation is safe.", AMBER, AMBER_BG),
    ("Translation must be correct", "A translator that emits plausible but wrong C++ is worse than none. Output must "
     "preserve meaning.", RED, RED_BG),
]
for i, (t, d, c, bgc) in enumerate(cards):
    x = 0.6 + i * 4.1
    rect(s, x, 1.75, 3.85, 2.75, fill=bgc, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    rect(s, x, 1.75, 0.09, 2.75, fill=c)
    text(s, x + 0.3, 1.95, 3.4, 0.5, t, size=19, color=c, bold=True, font=FONT_SB)
    text(s, x + 0.3, 2.65, 3.35, 1.8, d, size=16, color=INK)
rect(s, 0.6, 4.85, 12.13, 1.85, fill=NAVY, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
text(s, 0.95, 5.0, 11.5, 0.4, "PROBLEM STATEMENT", size=12, color=TEAL_L, bold=True)
text(s, 0.95, 5.38, 11.5, 1.25, "Build a compiler that accepts a simple pseudocode language, reports lexical, syntax "
     "and semantic errors precisely, generates and optimizes intermediate code, and emits equivalent C++, with "
     "every phase visible through a web interface.", size=18, color=WHITE)
notes(s, "Three gaps motivate the project: phases are taught separately, optimizers are a black box, and a "
         "translator is only useful if its output is correct. The problem statement on the slide is the project's "
         "formal goal.")

# ================================================================== 4 SOLUTION & OBJECTIVES
s = content_slide(4, "Proposed solution and objectives", "Solution")
rect(s, 0.6, 1.7, 5.4, 5.0, fill=LIGHT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
text(s, 0.9, 1.9, 4.9, 0.45, "The solution", size=20, color=TEAL, bold=True, font=FONT_SB)
bullets(s, 0.9, 2.45, 4.9, 4.1, [
    "A complete source-to-source compiler from pseudocode to C++17.",
    "Exposed as a JSON API: `POST /compile` returns every artifact.",
    "A browser workbench that renders each phase for the current program.",
    "Testing against reference interpreters and real g++ binaries.",
], size=16, gap=10)
objs = [
    "Design a small, complete language: int, float, string, bool, arithmetic, logic, print, if/else, while",
    "Implement every phase by hand, each testable without the web framework",
    "Report structured errors: stage, code, message, line and column",
    "Generate TAC, optimize it safely, and emit behaviour-preserving C++",
    "Present all phases in a responsive web app and deploy it publicly",
]
for i, o in enumerate(objs):
    y = 1.75 + i * 0.98
    rect(s, 6.4, y, 0.62, 0.62, fill=TEAL, line=None, shape=MSO_SHAPE.OVAL)
    text(s, 6.4, y + 0.08, 0.62, 0.45, str(i + 1), size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    text(s, 7.2, y - 0.02, 5.5, 0.85, o, size=16, color=INK, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "The solution has three parts: the compiler, the API and the workbench, tied together by testing. Read "
         "the five objectives; each is revisited later with evidence.")

# ================================================================== 5 FEATURES & LANGUAGE
s = content_slide(5, "Features and supported language constructs", "Language")
code_panel(s, 0.6, 1.7, 5.6, 4.95, [
    "# Classic FizzBuzz from 1 to 15",
    "let i = 1",
    "while i <= 15 {",
    "    if i % 15 == 0 {",
    '        print("FizzBuzz")',
    "    } else if i % 3 == 0 {",
    '        print("Fizz")',
    "    } else if i % 5 == 0 {",
    '        print("Buzz")',
    "    } else {",
    "        print(i)",
    "    }",
    "    i = i + 1",
    "}",
], title="examples/fizzbuzz.cb", size=14)
table(s, 6.6, 1.7, 6.13, [
    ["Construct", "Syntax"],
    ["Declaration", "`let x = 5`   `let f: float = 2`"],
    ["Assignment", "`x = x + 1`"],
    ["Output", "`print(a, b)`  or  `print a, b`"],
    ["Selection", "`if` … `else if` … `else` { }"],
    ["Iteration", "`while cond { }`"],
    ["Operators", "`+ - * / %`  `== != < <= > >=`  `and or not`"],
    ["Types → C++", "int→long long, float→double, string, bool"],
], [1.65, 4.48], size=13, row_h=0.43)
bullets(s, 6.6, 5.35, 6.1, 1.4, [
    "C++ semantics: `/` truncates, `%` follows the dividend, `and`/`or` short-circuit.",
    "Only int → float converts implicitly; conditions must be bool.",
], size=14, gap=4)
notes(s, "Walk through FizzBuzz: let, while, an else-if chain, print. The table lists everything the language "
         "supports; functions, arrays, for loops and input are intentionally out of scope.")

# ================================================================== 6 TECH STACK
s = content_slide(6, "Technology stack", "Tools")
cols = [
    ("Compiler + API", TEAL, TEAL_BG, ["Python 3.13.15", "FastAPI 0.143", "Pydantic 2.14", "Uvicorn 0.54"],
     "hand-written compiler package and JSON API"),
    ("Frontend", BLUE, BLUE_BG, ["React 19.3", "TypeScript 7.0", "Vite 8.3", "Tailwind CSS 4.3", "CodeMirror 6"],
     "workbench UI, editor and artifact views"),
    ("Testing", AMBER, AMBER_BG, ["pytest 9.1 + httpx2", "Vitest 5.0", "g++ 16.1 (C++17)", "GitHub Actions CI"],
     "unit, API, end-to-end and g++ build-and-run"),
    ("Deployment", RED, RED_BG, ["Render (API)", "render.yaml", "Vercel (frontend)", "vercel.json"],
     "both redeploy on every push to main"),
]
for i, (t, c, bgc, items, role) in enumerate(cols):
    x = 0.6 + i * 3.08
    rect(s, x, 1.75, 2.88, 4.25, fill=bgc, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    rect(s, x, 1.75, 2.88, 0.7, fill=c, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.25)
    text(s, x, 1.83, 2.88, 0.55, t, size=19, color=WHITE, bold=True, font=FONT_SB, align=PP_ALIGN.CENTER)
    text(s, x + 0.3, 2.7, 2.4, 3.2, [[("•  ", {"color": c, "bold": True}), (it, {"color": INK})] for it in items],
         size=17, after=10)
    rect(s, x + 0.25, 5.05, 2.38, 0.012, fill=c)
    text(s, x + 0.2, 5.15, 2.48, 0.8, role, size=13, color=SLATE, align=PP_ALIGN.CENTER)
text(s, 0.6, 6.25, 12.1, 0.6, [[("No parser generators or transpilers: ", {"bold": True, "color": TEAL}),
                                ("the lexer, parser, type checker, IR, optimizer and code generator are all original "
                                 "Python code.", {"color": INK})]], size=16)
notes(s, "Versions are the ones pinned in requirements.txt and package.json. The compiler uses only the Python "
         "standard library; FastAPI and Pydantic are only in the web layer.")

# ================================================================== 7 ARCHITECTURE
s = content_slide(7, "System architecture", "Architecture")
rect(s, 0.6, 1.65, 4.3, 5.1, fill=WHITE, line=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05, line_w=2)
text(s, 0.6, 1.75, 4.3, 0.4, "React frontend  ·  Vercel", size=17, color=BLUE, bold=True, font=FONT_SB,
     align=PP_ALIGN.CENTER)
box(s, 0.9, 2.3, 3.7, 1.2, "lib/api.ts · useCompiler", "POST /compile · GET /health\n650 ms auto-compile debounce",
    fill=BLUE_BG, edge=BLUE, title_size=15, sub_size=12)
box(s, 0.9, 3.75, 3.7, 1.15, "CodeMirror 6 editor", "highlighting · error underlines", fill=BLUE_BG, edge=BLUE,
    title_size=15, sub_size=12)
box(s, 0.9, 5.15, 3.7, 1.3, "Pipeline strip + 7 tabs", "C++ · Tokens · AST · Symbols\nTAC · Optimized TAC · Diagnostics",
    fill=BLUE_BG, edge=BLUE, title_size=15, sub_size=12)
rect(s, 6.15, 1.65, 6.58, 5.1, fill=WHITE, line=TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05, line_w=2)
text(s, 6.15, 1.75, 6.58, 0.4, "FastAPI backend  ·  Render", size=17, color=TEAL, bold=True, font=FONT_SB,
     align=PP_ALIGN.CENTER)
box(s, 6.45, 2.3, 2.95, 1.2, "app/main.py", "GET /health · POST /compile\nCORS · 413 · JSON 500", title_size=15,
    sub_size=12)
box(s, 9.55, 2.3, 2.9, 1.2, "config · schemas", "CORS_ORIGINS\nMAX_SOURCE_CHARS · Pydantic", title_size=15, sub_size=12)
box(s, 6.45, 3.75, 6.0, 0.9, "pipeline.compile_source()", "runs phases in order, times each, stops at first failure",
    title_size=15, sub_size=12)
mods = ["lexer", "parser", "semantic", "tacgen", "optimizer", "codegen"]
for i, m in enumerate(mods):
    box(s, 6.45 + i * 1.0, 5.15, 0.93, 0.75, m, None, fill=LIGHT, edge=TEAL, title_size=11)
text(s, 6.45, 6.0, 6.0, 0.35, "app/compiler: no web dependencies, each module unit-tested", size=12, color=SLATE,
     align=PP_ALIGN.CENTER)
line(s, 4.6, 2.65, 6.45, 2.65, color=BLUE, width=2)
line(s, 6.45, 3.15, 4.6, 3.15, color=TEAL, width=2)
text(s, 4.75, 2.25, 1.6, 0.35, "JSON request", size=12, color=BLUE, align=PP_ALIGN.CENTER)
text(s, 4.75, 3.2, 1.6, 0.35, "JSON response", size=12, color=TEAL, align=PP_ALIGN.CENTER)
line(s, 7.9, 3.5, 7.9, 3.75, width=1.75)
line(s, 9.45, 4.65, 9.45, 5.15, width=1.75)
notes(s, "Two tiers. The browser sends the editor contents as JSON to POST /compile. main.py validates the request, "
         "applies the size limit and CORS, and calls compile_source, which runs the six compiler modules. The compiler "
         "package has no FastAPI imports, so every phase is unit-tested on its own. VITE_API_URL is baked into the "
         "frontend at build time; CORS on Render allows only the Vercel origin.")

# ================================================================== 8 PIPELINE
s = content_slide(8, "Compiler pipeline: six phases", "Pipeline")
phases = [("Lexer", "lexer.py", "tokens"), ("Parser", "parser.py", "AST"), ("Semantic", "semantic.py", "typed AST\n+ symbols"),
          ("TAC", "tacgen.py", "three-address\ncode"), ("Optimizer", "optimizer.py", "optimized\nTAC"),
          ("C++ gen", "codegen_cpp.py", "C++17\nsource")]
for i, (t, m, o) in enumerate(phases):
    x = 0.6 + i * 2.06
    sh = rect(s, x, 1.85, 1.98, 1.15, fill=TEAL if i % 2 == 0 else RGBColor(0x0F, 0x76, 0x6E), line=None,
              shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
    sh.adjustments[0] = 0.28
    text(s, x + 0.25, 1.98, 1.55, 0.9, [[(t, {"bold": True, "size": 17, "color": WHITE, "font": FONT_SB})],
                                         [(m, {"size": 11, "color": RGBColor(0xCC, 0xFB, 0xF1)})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=0)
    text(s, x + 0.1, 3.1, 1.8, 0.7, o, size=13, color=INK, align=PP_ALIGN.CENTER, spacing=1.0)
text(s, 0.6, 1.45, 3, 0.35, "source text ▸", size=12, color=SLATE)
for i, code in enumerate(["L001–L005", "P001–P002", "S001–S006"]):
    cx = 0.6 + i * 2.06 + 0.99
    line(s, cx, 3.85, cx, 4.35, color=RED, dash=True, width=1.5)
    text(s, cx - 0.9, 3.85, 0.85, 0.35, code, size=11, color=RED, align=PP_ALIGN.RIGHT)
rect(s, 0.6, 4.4, 6.2, 0.65, fill=RED_BG, line=RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
text(s, 0.6, 4.47, 6.2, 0.5, "Structured diagnostics: stage · code · message · line:column · hints", size=14,
     color=RED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 7.1, 4.35, 5.6, 0.8, "A failing phase stops the pipeline: later phases are reported as skipped and no TAC "
     "or C++ is returned.", size=14, color=SLATE)
table(s, 0.6, 5.35, 12.13, [
    ["Artifact in the API response", "tokens", "ast", "symbols", "tac", "optimized_tac", "cpp"],
    ["Workbench tab", "Tokens", "AST", "Symbols", "TAC", "Optimized TAC", "C++"],
], [3.33, 1.4, 1.4, 1.5, 1.4, 1.75, 1.35], size=13, row_h=0.45)
notes(s, "The six phases, the module that implements each and what it produces. Only the first three phases can "
         "report user errors: lexical L001 to L005, syntax P001 and P002, semantic S001 to S006. When a phase fails "
         "the pipeline stops, so the UI never shows TAC or C++ for an invalid program. Each artifact maps to one tab.")

# ================================================================== 9 LEXER / PARSER / AST
s = content_slide(9, "Lexer, parser and abstract syntax tree", "Front end of the compiler")
text(s, 0.6, 1.6, 6.3, 0.45, "Lexical analysis  (lexer.py)", size=19, color=TEAL, bold=True, font=FONT_SB)
bullets(s, 0.6, 2.05, 6.3, 1.9, [
    "Hand-written scanner; longest-match operators (`<=` before `<`).",
    "1-based line:column span on every token; newlines ignored inside ( ).",
    "Reports **all** lexical errors in one pass, with hints (`!` → use `not`).",
], size=15, gap=5)
text(s, 0.6, 3.85, 6.3, 0.45, "Syntax analysis  (parser.py)", size=19, color=TEAL, bold=True, font=FONT_SB)
bullets(s, 0.6, 4.3, 6.3, 2.6, [
    "Recursive descent: one method per grammar rule.",
    "Precedence from the call chain: `or` < `and` < `not` < equality < comparison < `+ -` < `* / %` < unary.",
    "Typed AST dataclasses with spans, serialized to JSON for the AST tab.",
    "Stops at the first syntax error; nesting ≤ 50, expression height ≤ 150.",
], size=15, gap=5)
picture(s, CROPS / "ast-panel.png", 7.25, 1.65, w=5.48)
text(s, 7.25, 5.35, 5.48, 0.6, "Live AST tab: typed nodes for an else-if chain (cropped screenshot)", size=11,
     color=SLATE, align=PP_ALIGN.CENTER)
notes(s, "The lexer collects every lexical error so the user sees them all at once. The parser is predictive "
         "recursive descent; precedence is encoded by which method calls which, and left associativity by loops. "
         "Depth limits stop malicious input from overflowing Python's call stack. The screenshot is the live AST tab.")

# ================================================================== 10 SEMANTIC / TAC / OPTIMIZATION
s = content_slide(10, "Semantic analysis, intermediate code and optimization", "Middle end")
text(s, 0.6, 1.6, 4.0, 0.45, "Semantic analysis", size=18, color=TEAL, bold=True, font=FONT_SB)
bullets(s, 0.6, 2.05, 4.0, 4.8, [
    "Symbol table = stack of scopes, one per block.",
    "S001 undefined · S002 redeclared · S003 bad assignment · S004 operand types · S005 non-bool condition · "
    "S006 division by literal 0.",
    "Shadowed names renamed (`x` → `x_2`).",
    "All errors collected, no cascades; S100 unused warning.",
], size=14, gap=6)
code_panel(s, 4.85, 1.65, 3.85, 2.75, ["t1 = 2 * 3", "t2 = t1 + 4", "x = t2", "i = 0", "L1:", "t3 = i < x",
                                       "ifFalse t3 goto L2"], title="Original TAC", size=13)
code_panel(s, 4.85, 4.55, 3.85, 2.25, ["x = 10", "i = 0", "L1:", "t3 = i < x", "ifFalse t3 goto L2"],
           title="Optimized TAC", size=13)
text(s, 9.0, 1.6, 3.8, 0.45, "Optimizer: 11 passes", size=18, color=TEAL, bold=True, font=FONT_SB)
text(s, 9.0, 2.02, 3.8, 0.3, "repeated until nothing changes", size=12, color=SLATE)
passes = ["Constant folding", "Constant propagation", "Copy propagation", "Algebraic simplification",
          "Self-assignment removal", "Branch folding", "Temporary coalescing", "Dead temporary elimination",
          "Unreachable code removal", "Redundant jump removal", "Unused label removal"]
text(s, 9.0, 2.4, 3.75, 3.3, [[("✓  ", {"color": TEAL, "bold": True}), (p, {"color": INK})] for p in passes],
     size=13, after=2)
rect(s, 9.0, 5.75, 3.73, 1.05, fill=TEAL_BG, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
text(s, 9.1, 5.8, 3.55, 0.95, [[("Safe by design: ", {"bold": True, "color": TEAL}),
                                ("local to basic blocks, no overflow or ÷0 folding, exact float identities only.",
                                 {"color": INK})]], size=12.5, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Semantic analysis assigns a type to every expression and reports the six semantic error codes. TAC is "
         "the intermediate representation: one operator per instruction, labels and conditional jumps. The code panels "
         "show exact compiler output for let x = 2 * 3 + 4 before a loop: folding and propagation reduce three "
         "instructions to x = 10. Propagation is reset at labels, so it never crosses into a loop body.")

# ================================================================== 11 CODEGEN / ERRORS / OUTPUT
s = content_slide(11, "C++ code generation, error handling and output", "Back end")
code_panel(s, 0.6, 1.65, 4.05, 3.65, ["let n = 1", "let total = 0", "while n <= 5 {", "  if n % 2 == 0 {",
                                     "    total = total + n * n", "  } else {", '    print n, "is odd"', "  }",
                                     "  n = n + 1", "}", 'print("sum of even squares:", total)'],
           title="CodeBridge source", size=12)
code_panel(s, 4.8, 1.65, 4.6, 3.65, ["long long n = 1;", "long long total = 0;", "while (n <= 5) {",
                                       "  if (n % 2 == 0) {", "    total = total + n * n;", "  } else {",
                                       '    std::cout << n << " is odd\\n";', "  }", "  n = n + 1;", "}",
                                       'std::cout << "sum of even squares: "', '          << total << "\\n";'],
           title="Generated C++ (body of main)", size=12)
bullets(s, 0.6, 5.45, 8.8, 1.5, [
    "Built from **optimized TAC**: single-use temporaries inlined, jump patterns rebuilt as `while` / `if` / `else if`, "
    "`goto` fallback.",
    "g++ output: `1 is odd`, `3 is odd`, `5 is odd`, `sum of even squares: 20`: identical to the reference interpreter.",
], size=14, gap=4)
picture(s, CROPS / "diagnostics-panel.png", 9.6, 1.65, w=3.13)
text(s, 9.6, 3.97, 3.13, 0.55, "Live Diagnostics tab (cropped): code, line:col, message, hint", size=11,
     color=SLATE, align=PP_ALIGN.CENTER, spacing=1.0)
bullets(s, 9.6, 4.65, 3.15, 2.3, [
    "Errors underlined in the editor",
    "413 if too large, 422 if malformed",
    "JSON 500, never a crash",
    "No C++ after a failure",
], size=12.5, gap=3)
notes(s, "This pair is unedited compiler output; the last C++ statement is wrapped here only to fit the slide. The "
         "generator never sees the AST, only optimized TAC, and reconstructs structured control flow from jump "
         "patterns. On the right is the error handling: structured diagnostics, HTTP limits, and the rule that no "
         "C++ is produced for a failing program.")

# ================================================================== 12 SCREENSHOTS & WORKFLOW
s = content_slide(12, "Live application and user workflow", "Demo")
picture(s, SHOTS / "live-optimized.png", 0.6, 1.6, w=7.6)
text(s, 0.6, 6.4, 7.6, 0.35, "codebridge-swart.vercel.app: Optimizer showcase, 25 → 14 TAC instructions (−44%)",
     size=11.5, color=SLATE, align=PP_ALIGN.CENTER)
steps = [
    ("Open", "codebridge-swart.vercel.app; first request may wake the free-tier server"),
    ("Write or pick", "type a program or choose one of 12 examples"),
    ("Compile", "auto-compiles 650 ms after typing, or press Compile (Ctrl+Enter)"),
    ("Inspect", "switch tabs: Tokens, AST, Symbols, TAC, Optimized TAC"),
    ("Use the C++", "copy the generated C++ and build it with any C++17 compiler"),
]
for i, (t, d) in enumerate(steps):
    y = 1.6 + i * 0.97
    rect(s, 8.55, y, 0.55, 0.55, fill=TEAL, line=None, shape=MSO_SHAPE.OVAL)
    text(s, 8.55, y + 0.07, 0.55, 0.4, str(i + 1), size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    text(s, 9.25, y - 0.05, 3.5, 0.9, [[(t, {"bold": True, "color": INK, "size": 15})],
                                        [(d, {"color": SLATE, "size": 12})]], after=0, spacing=1.0)
notes(s, "This is a screenshot of the deployed site, not a mock-up. If the network allows, switch to the live "
         "site and follow the five steps: load the Optimizer showcase, show Optimized TAC, introduce a typo to show an "
         "inline error, then open the C++ tab.")

# ================================================================== 13 TESTING
s = content_slide(13, "Testing strategy and verified results", "Verification")
table(s, 0.6, 1.6, 7.3, [
    ["Automated suite", "Tests", "Highlights"],
    ["Lexer", "12", "every lexical error code"],
    ["Parser + AST", "48", "precedence, 21 syntax errors"],
    ["Semantic", "31", "S001–S006, scopes"],
    ["TAC", "18", "TAC vs. AST interpreter"],
    ["Optimizer", "28", "300 random programs"],
    ["C++ generation", "36", "g++ build and run"],
    ["API + health", "34", "413/422/500, CORS"],
    ["End-to-end", "6", "server → HTTP → g++"],
    ["Frontend (Vitest)", "20", "API client, views"],
], [2.35, 0.95, 4.0], size=13, row_h=0.42)
text(s, 0.6, 5.95, 7.3, 0.9, [[("Oracle: ", {"bold": True, "color": TEAL}),
                               ("each example must print the same output in the AST interpreter, original TAC, "
                                "optimized TAC and a g++-built binary.", {"color": INK})]], size=13.5)
stat(s, 8.25, 1.6, 2.15, "213", "backend tests\npassed")
stat(s, 10.58, 1.6, 2.15, "20", "frontend tests\npassed")
rect(s, 8.25, 3.3, 4.48, 3.5, fill=LIGHT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
text(s, 8.45, 3.4, 4.1, 0.4, "Manual checks (not automated)", size=15, color=TEAL, bold=True, font=FONT_SB)
bullets(s, 8.45, 3.85, 4.15, 2.95, [
    "`/health` → HTTP 200",
    "CORS: Vercel origin allowed, others rejected",
    "Headless Edge: app online, all phases green",
    "Live C++ for 3 examples built with g++ = expected output",
    "Warm compile round trip 0.55–1.45 s",
    "Audit: 192 random programs matched g++",
], size=12.5, gap=3)
notes(s, "Distinguish the two kinds of evidence. The table is the automated suite, run locally and in GitHub "
         "Actions: 213 backend and 20 frontend tests, all passing. The box on the right lists checks done by hand or "
         "with one-off scripts against the deployed site on 13 October 2026; they are not part of the automated suite.")

# ================================================================== 14 LIMITATIONS & FUTURE
s = content_slide(14, "Limitations and future enhancements", "Next steps")
for col, (title, c, bgc, items) in enumerate([
    ("Current limitations", AMBER, AMBER_BG, [
        "No functions, arrays, `for` loops or input",
        "Only the first syntax error is reported",
        "Optimization is local to basic blocks; no loop-invariant code motion",
        "Run-time ÷0 or overflow is undefined, as in C++",
        "Free hosting sleeps when idle: cold start up to ~1 minute",
    ]),
    ("Future enhancements", TEAL, TEAL_BG, [
        "Functions with a call stack in TAC",
        "Panic-mode error recovery for multiple syntax errors",
        "Global data-flow analysis and SSA form",
        "Run generated C++ in the browser via WebAssembly",
    ]),
]):
    x = 0.6 + col * 6.2
    rect(s, x, 1.7, 5.93, 5.1, fill=bgc, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    rect(s, x, 1.7, 5.93, 0.75, fill=c, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
    text(s, x, 1.8, 5.93, 0.55, title, size=21, color=WHITE, bold=True, font=FONT_SB, align=PP_ALIGN.CENTER)
    bullets(s, x + 0.35, 2.7, 5.3, 4.0, items, size=17, gap=12, marker_color=c)
notes(s, "Be candid: the language is intentionally small and the optimizer is intentionally conservative. The "
         "future work list follows from the limitations; WebAssembly would show program output without executing "
         "code on the server.")

# ================================================================== 15 CONCLUSION & Q&A
s = dark_slide()
rect(s, 0, 0, 13.333, 0.08, fill=TEAL)
text(s, 0.75, 0.6, 7, 0.4, "CONCLUSION", size=14, color=TEAL_L, bold=True, font=FONT_SB)
text(s, 0.75, 1.0, 7.4, 0.8, "What CodeBridge delivers", size=32, color=WHITE, bold=True, font=FONT_SB)
bullets(s, 0.75, 1.95, 7.2, 3.8, [
    "A complete, hand-written six-phase compiler from pseudocode to C++17.",
    "Every phase visible in a deployed, responsive web workbench.",
    "Safe, measurable optimization (e.g. 25 → 14 TAC instructions).",
    "Correctness backed by 233 automated tests, reference interpreters and real g++ builds.",
], size=18, color=WHITE, gap=12, marker_color=TEAL_L)
rect(s, 8.55, 1.0, 4.2, 4.7, fill=NAVY2, line=RGBColor(0x33, 0x41, 0x55), shape=MSO_SHAPE.ROUNDED_RECTANGLE,
     radius=0.06)
text(s, 8.55, 1.55, 4.2, 1.0, "Thank you", size=34, color=WHITE, bold=True, font=FONT_SB, align=PP_ALIGN.CENTER)
text(s, 8.55, 2.5, 4.2, 0.8, "Questions?", size=28, color=TEAL_L, font=FONT_SB, align=PP_ALIGN.CENTER)
text(s, 8.8, 3.6, 3.7, 1.8, [
    [("Live app", {"bold": True, "color": TEAL_L, "size": 13})],
    [("codebridge-swart.vercel.app", {"color": WHITE, "size": 14})],
    [("Source code", {"bold": True, "color": TEAL_L, "size": 13})],
    [("github.com/ashmita-dev/CodeBridge", {"color": WHITE, "size": 14})],
], align=PP_ALIGN.CENTER, after=3)
text(s, 0.75, 6.4, 12, 0.4, "[Student Name]  ·  [Roll Number]  ·  Guide: [Guide Name]  ·  [Institution]", size=13,
     color=RGBColor(0x94, 0xA3, 0xB8))
notes(s, "Summarize the four outcomes, then invite questions. Likely questions: why recursive descent (simple, "
         "precise errors, no generator), why optimize TAC rather than the AST (classical, easier to reason about), "
         "how correctness is checked (reference interpreters and g++), and why code is not executed on the server "
         "(security).")

assert len(prs.slides) == TOTAL, len(prs.slides)
prs.core_properties.title = "CodeBridge: Pseudocode-to-C++ Compiler"
prs.core_properties.author = "[Student Name]"
prs.save(OUT)
print("wrote", OUT, "with", len(prs.slides), "slides")
